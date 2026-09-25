"""
Script to add Cabinet Falcon as the host company, rewrite the contextualisation
to connect Falcon to the predictive maintenance project, and renumber sections.
"""
from docx import Document
from docx.shared import Pt
from copy import deepcopy

INPUT_PATH  = r'C:\Users\ULTRAPC\Desktop\pfa2026\Rapport_PFA2026_updated.docx'
OUTPUT_PATH = r'C:\Users\ULTRAPC\Desktop\pfa2026\Rapport_PFA2026_v2.docx'

doc = Document(INPUT_PATH)

# ─────────────── Helper ───────────────
def insert_paragraph_after(paragraph, text, style='Normal'):
    """Insert a new paragraph after the given one."""
    new_p = deepcopy(paragraph._element)
    for child in list(new_p):
        if child.tag.endswith('}r'):
            new_p.remove(child)
    paragraph._element.addnext(new_p)
    from docx.text.paragraph import Paragraph
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = doc.styles[style]
    new_para.clear()
    new_para.add_run(text)
    return new_para

# ─────────────── Locate key paragraphs ───────────────
indices = {}
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t == "1.2.4 Contextualisation du projet":
        indices['ctx_heading'] = i
    elif t.startswith("Le pr\u00e9sent projet de fin d'ann\u00e9e s'inscrit"):
        indices['ctx_p1'] = i
    elif t.startswith("Le projet s'appuie sur le jeu de donn\u00e9es NASA"):
        indices['ctx_p2'] = i
    elif t.startswith("La solution d\u00e9velopp\u00e9e comprend un pipeline"):
        indices['ctx_p3'] = i
    elif t.startswith("Une attention particuli\u00e8re a \u00e9t\u00e9 accord\u00e9e"):
        indices['ctx_p4'] = i
    elif t == "1.3 Probl\u00e9matique":
        indices['prob_heading'] = i
    elif t.startswith("L'observation de la litt\u00e9rature"):
        indices['prob_p1'] = i
    elif t.startswith("D\u00e8s lors, notre probl\u00e9matique"):
        indices['prob_p2'] = i
    elif t == "1.4 Objectifs du projet":
        indices['obj_heading'] = i
    elif t == "1.5 Conclusion":
        indices['concl_heading'] = i

print("Found indices:", indices)

# ═══════════════════════════════════════════════════════════════════════
# STEP 1: Replace "1.2.4 Contextualisation du projet" (Heading 3)
#         with "1.3 Présentation de l'organisme d'accueil" (Heading 2)
#         and rewrite its content to be about Cabinet Falcon
# ═══════════════════════════════════════════════════════════════════════

# Change the heading
doc.paragraphs[indices['ctx_heading']].clear()
doc.paragraphs[indices['ctx_heading']].add_run("1.3 Pr\u00e9sentation de l\u2019organisme d\u2019accueil")
doc.paragraphs[indices['ctx_heading']].style = doc.styles['Heading 2']

# Replace the 4 context paragraphs with Falcon presentation
falcon_paragraphs = [
    (
        "Mon projet de fin d\u2019ann\u00e9e s\u2019est d\u00e9roul\u00e9 au sein du cabinet Falcon, "
        "un acteur sp\u00e9cialis\u00e9 dans le contr\u00f4le qualit\u00e9, l\u2019audit, l\u2019inspection "
        "et l\u2019accompagnement \u00e0 la certification, bas\u00e9 \u00e0 Casablanca. Cr\u00e9\u00e9 en 2010, "
        "Falcon s\u2019est progressivement impos\u00e9 comme un acteur de r\u00e9f\u00e9rence dans "
        "son segment au niveau r\u00e9gional, en d\u00e9veloppant une expertise reconnue "
        "notamment dans l\u2019accompagnement des entreprises \u00e0 la certification "
        "ISO 9001:2015 et ISO 14001:2015."
    ),
    (
        "Fort d\u2019une \u00e9quipe de 15 consultants qualit\u00e9 et 3 auditeurs internes, "
        "le cabinet r\u00e9alise environ 20 \u00e0 25 missions par an pour un portefeuille "
        "de clients allant des PME aux grands groupes industriels, r\u00e9partis dans "
        "plusieurs secteurs d\u2019activit\u00e9 : industrie manufacturi\u00e8re, BTP, services, "
        "agroalimentaire et logistique. Son chiffre d\u2019affaires annuel est estim\u00e9 "
        "\u00e0 environ 4,5 millions de MAD. L\u2019organisation repose sur une structure "
        "fonctionnelle classique articul\u00e9e autour de quatre d\u00e9partements : "
        "Commercial, Technique, Contr\u00f4le de Gestion & Administration, et le "
        "Service Comptabilit\u00e9 & Facturation."
    ),
    (
        "Affect\u00e9 au sein du D\u00e9partement Contr\u00f4le de Gestion & Administration, "
        "j\u2019ai b\u00e9n\u00e9fici\u00e9 d\u2019un acc\u00e8s privil\u00e9gi\u00e9 aux outils de pilotage, aux "
        "donn\u00e9es op\u00e9rationnelles des \u00e9quipements suivis par le cabinet et aux "
        "processus de maintenance des clients industriels. C\u2019est dans ce "
        "contexte que la probl\u00e9matique de la maintenance pr\u00e9dictive a \u00e9merg\u00e9 : "
        "les entreprises clientes de Falcon, en particulier dans les secteurs "
        "manufacturier et logistique, font face \u00e0 des arr\u00eats non planifi\u00e9s "
        "co\u00fbteux que les strat\u00e9gies de maintenance corrective et pr\u00e9ventive "
        "ne parviennent pas \u00e0 \u00e9liminer."
    ),
    (
        "Le projet a ainsi consist\u00e9 \u00e0 concevoir un syst\u00e8me intelligent de "
        "maintenance pr\u00e9dictive exploitant l\u2019intelligence artificielle pour "
        "estimer la dur\u00e9e de vie utile restante (Remaining Useful Life \u2013 RUL) "
        "des \u00e9quipements industriels. La solution s\u2019appuie sur le jeu de donn\u00e9es "
        "NASA C-MAPSS FD001, un r\u00e9f\u00e9rentiel reconnu dans la communaut\u00e9 "
        "scientifique, et int\u00e8gre un pipeline d\u2019apprentissage automatique, "
        "un backend REST d\u00e9velopp\u00e9 avec FastAPI, une base de donn\u00e9es "
        "relationnelle PostgreSQL et une interface web de supervision "
        "d\u00e9velopp\u00e9e avec React et TypeScript."
    ),
]

ctx_indices = [indices['ctx_p1'], indices['ctx_p2'], indices['ctx_p3'], indices['ctx_p4']]
for idx, new_text in zip(ctx_indices, falcon_paragraphs):
    doc.paragraphs[idx].clear()
    doc.paragraphs[idx].add_run(new_text)

print("\u2713 Section 1.3 - Falcon presentation written")

# ═══════════════════════════════════════════════════════════════════════
# STEP 2: Renumber 1.3 Problématique -> 1.4 Problématique
# ═══════════════════════════════════════════════════════════════════════
doc.paragraphs[indices['prob_heading']].clear()
doc.paragraphs[indices['prob_heading']].add_run("1.4 Probl\u00e9matique")
print("\u2713 1.3 -> 1.4 Probl\u00e9matique renumbered")

# ═══════════════════════════════════════════════════════════════════════
# STEP 3: Renumber 1.4 Objectifs -> 1.5 Objectifs du projet
# ═══════════════════════════════════════════════════════════════════════
doc.paragraphs[indices['obj_heading']].clear()
doc.paragraphs[indices['obj_heading']].add_run("1.5 Objectifs du projet")
print("\u2713 1.4 -> 1.5 Objectifs renumbered")

# ═══════════════════════════════════════════════════════════════════════
# STEP 4: Renumber 1.5 Conclusion -> 1.6 Conclusion
# ═══════════════════════════════════════════════════════════════════════
doc.paragraphs[indices['concl_heading']].clear()
doc.paragraphs[indices['concl_heading']].add_run("1.6 Conclusion")
print("\u2713 1.5 -> 1.6 Conclusion renumbered")

# ═══════════════════════════════════════════════════════════════════════
# STEP 5: Update the Problématique to connect with Falcon
# ═══════════════════════════════════════════════════════════════════════
enriched_prob = (
    "L\u2019observation du terrain au sein du cabinet Falcon et l\u2019analyse des "
    "besoins de ses clients industriels r\u00e9v\u00e8lent un paradoxe : bien que les "
    "entreprises disposent de volumes croissants de donn\u00e9es issues de leurs "
    "capteurs industriels, la transformation de ces donn\u00e9es en d\u00e9cisions de "
    "maintenance exploitables reste un d\u00e9fi majeur. Les strat\u00e9gies de "
    "maintenance corrective et pr\u00e9ventive, encore largement dominantes, "
    "engendrent soit des arr\u00eats non planifi\u00e9s co\u00fbteux, soit des interventions "
    "pr\u00e9matur\u00e9es sur des \u00e9quipements encore fonctionnels. Par ailleurs, les "
    "mod\u00e8les pr\u00e9dictifs, lorsqu\u2019ils ne font pas l\u2019objet d\u2019une validation "
    "rigoureuse, peuvent produire des r\u00e9sultats artificiellement performants "
    "en raison de fuites de donn\u00e9es, de variables confondantes ou de "
    "protocoles d\u2019\u00e9valuation inadapt\u00e9s aux s\u00e9ries temporelles."
)

doc.paragraphs[indices['prob_p1']].clear()
doc.paragraphs[indices['prob_p1']].add_run(enriched_prob)
print("\u2713 Probl\u00e9matique body updated with Falcon context")

# ── Save ──
doc.save(OUTPUT_PATH)
print(f"\n\u2705 Saved to: {OUTPUT_PATH}")
